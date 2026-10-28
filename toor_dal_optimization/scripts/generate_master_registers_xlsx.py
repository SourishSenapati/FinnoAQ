# -*- coding: utf-8 -*-
"""
Generate Master Governance Excel Workbooks:
1. docs/evidence/EVIDENCE_MATRIX.xlsx
2. docs/experiments/EXPERIMENT_REGISTER.xlsx
3. docs/acceptance/DELIVERABLE_ACCEPTANCE_CHECKLIST.xlsx
"""

import os
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

# Style Constants
NAVY_FILL = PatternFill(start_color="1A365D", end_color="1A365D", fill_type="solid")
STEEL_FILL = PatternFill(start_color="2B6CB0", end_color="2B6CB0", fill_type="solid")
ALT_FILL = PatternFill(start_color="F7FAFC", end_color="F7FAFC", fill_type="solid")
WHITE_FONT_BOLD = Font(name="Calibri", size=10, bold=True, color="FFFFFF")
TITLE_FONT = Font(name="Calibri", size=12, bold=True, color="1A365D")
SUBTITLE_FONT = Font(name="Calibri", size=9, italic=True, color="4A5568")
DATA_FONT = Font(name="Calibri", size=9, color="2D3748")
DATA_FONT_BOLD = Font(name="Calibri", size=9, bold=True, color="1A365D")

THIN_BORDER = Border(
    left=Side(style='thin', color='CBD5E0'),
    right=Side(style='thin', color='CBD5E0'),
    top=Side(style='thin', color='CBD5E0'),
    bottom=Side(style='thin', color='CBD5E0')
)

def auto_fit_columns(ws, max_len_cap=60):
    for col in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            val_str = str(cell.value or '')
            if '\n' in val_str:
                val_str = max(val_str.split('\n'), key=len)
            if len(val_str) > max_len:
                max_len = len(val_str)
        ws.column_dimensions[col_letter].width = min(max(max_len + 3, 11), max_len_cap)

# ==============================================================================
# 1. EVIDENCE MATRIX WORKBOOK
# ==============================================================================
def create_evidence_matrix_wb(path):
    wb = openpyxl.Workbook()
    
    # Sheet 1: Master Claims
    ws1 = wb.active
    ws1.title = "Evidence_Matrix"
    ws1.views.sheetView[0].showGridLines = True
    
    ws1["A1"] = "FINNOAQ TOOR DAL R&D — MASTER CLAIM & EVIDENCE MATRIX"
    ws1["A1"].font = TITLE_FONT
    ws1["A2"] = "Project Ref: FNQ-PR-2026-001 | Governing Ref: FIN-GOV-LAW-002 | Date: 20-Oct-2026"
    ws1["A2"].font = SUBTITLE_FONT
    
    headers = [
        "Claim ID", "Claim Statement / Numerical Value", "Epistemic Category",
        "Primary Source Citation / File Location", "Inputs & Conventions",
        "Independent Verification / Formula", "Limitations & Epistemic Boundaries",
        "Engineering Disposition", "Deliverable Location"
    ]
    
    for col_idx, h in enumerate(headers, 1):
        cell = ws1.cell(row=4, column=col_idx, value=h)
        cell.fill = NAVY_FILL
        cell.font = WHITE_FONT_BOLD
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = THIN_BORDER
    ws1.row_dimensions[4].height = 28
    
    claims_data = [
        ("CLM-001", "Client preliminary observation of ~6 whistles (Tanzanian) vs ~3 whistles (Indian).",
         "SOURCE_REPORTED", "Phase1_Scope_of_Work_Toor_Dal.docx: Sec 1", "Domestic 5L pressure cooker, high heat.",
         "Unstandardized operator observation.", "Not an instrumented measurement; subject to fill ratio & heat flux.",
         "RETAIN_AS_CONTEXT", "Proposal Sec 2, Annex E"),
        ("CLM-002", "8192x8192 Mesh Grid Solve requires 1.50 GB VRAM & 30.75 ms/step (not 0.0051s).",
         "CALCULATED", "high_res_spatial_solver.py:141-167", "67.1M nodes, float32, 2D FD stencil.",
         "Profiled on RTX 4050 GPU via true stencil benchmark.", "0D Green's function shortcut discredited; real stencil verified.",
         "REVISE_AND_DISCREDIT", "FORENSIC_AUDIT_AND_CLAIM_LEDGER.md"),
        ("CLM-003", "Pectin de-esterification and calcium pectate crosslinking in middle lamella.",
         "HYPOTHESIS", "Literature (Jones & Boulter 1983; Hincks & Stanley 1986)", "Assumed monovalent Na+ displacement.",
         "Zero wet-lab assays conducted on target Tanzanian lot.", "Must compete with H1-H10; cannot be asserted without wet data.",
         "DEMOTE_TO_HYPOTHESIS", "Proposal Sec 6, Annex D"),
        ("CLM-004", "Incoming Raw Feed = 15,000 kg/day at 10.5% wet-basis moisture.",
         "ASSUMPTION", "Phase1_Scope_of_Work_Toor_Dal.docx / Design Basis", "15,000 kg/day feed, X_wb = 0.105",
         "Dry solids = 15,000 * 0.895 = 13,425.000 kg/day.", "Defines raw feed basis; distinct from dry dal entering soaking.",
         "RETAIN_WITH_BASIS", "Proposal Annex B, StreamBalance.xlsx"),
        ("CLM-005", "Absolute Dry Solids Loss during processing = 184.000 kg/day (1.3706%).",
         "ASSUMPTION", "Master Design Basis / Previous Notes", "184 kg solids lost in wash/peel.",
         "Net dry solids = 13,425 - 184 = 13,241.000 kg/day.", "Distinct from alternate 1.46% (196.005 kg); 13,241 kg adopted.",
         "RECONCILE_AND_ADOPT", "Proposal Annex B, MASS_AND_MOISTURE_BALANCE_AUDIT.md"),
        ("CLM-006", "Final Packaged Dal Moisture = 10.0% wet basis (14,712.222 kg/day product).",
         "OFFICIAL_SOURCE", "FSSAI Food Product Standards 2.4.6.14", "Max statutory moisture = 14.0%; target = 10.0%.",
         "Product = 13,241 / (1 - 0.10) = 14,712.222 kg/day.", "Guarantees storage stability and mold prevention.",
         "RETAIN_VERIFIED", "Proposal Annex B, Product Specs"),
        ("CLM-007", "Pre-Dryer Feed Mass = 18,390.278 kg/day under true 28.0% WB (Scenario B1).",
         "CALCULATED", "First-principles mass balance", "Dry solids = 13,241 kg, X_wb = 0.2800",
         "Wet mass = 13,241 / 0.72 = 18,390.278 kg/day.", "Water = 5,149.278 kg; Evap = 3,678.056 kg/day.",
         "RETAIN_AS_SCENARIO_B1", "Proposal Annex B, MASS_AND_MOISTURE_BALANCE_AUDIT.md"),
        ("CLM-008", "Pre-Dryer Feed Mass = 18,590.000 kg/day implies 28.7735% WB (Scenario B2).",
         "CALCULATED", "Previous draft stream table", "Wet mass = 18,590 kg, Dry solids = 13,241 kg",
         "Moisture = (18,590 - 13,241) / 18,590 = 28.7735% WB.", "Water = 5,349.000 kg; Evap = 3,877.778 kg/day (+199.7 kg delta).",
         "RETAIN_AS_SCENARIO_B2", "Proposal Annex B, MASS_AND_MOISTURE_BALANCE_AUDIT.md"),
        ("CLM-009", "Retained elemental sodium = 53.9 mg / 100g dal under single mist rinse.",
         "CALCULATED", "Stoichiometric washing balance", "0.50% NaHCO3 soak, 25% displacement rinse.",
         "Na element = 27.37% of NaHCO3; compliant with FSSAI.", "Requires wet-lab flame photometry / ICP-OES validation.",
         "RETAIN_AS_ESTIMATE", "Proposal Sec 11, Food Safety Matrix"),
        ("CLM-010", "Phase 1 Fixed Professional Fee = ₹1,73,000 + GST across 11 line items.",
         "OFFICIAL_SOURCE", "Phase1_Scope_of_Work_Toor_Dal.docx: Table 40", "Itemized tasks 01 to 11.",
         "Sum(Items 01..11) = 20k+7k+5k+20k+10k+18k+16k+17k+20k+20k+20k = ₹1,73,000.", "Exact 100.0% match; ceiling is ₹2,00,000 + GST.",
         "RETAIN_VERIFIED", "Proposal Sec 14, Commercial Basis Dossier"),
        ("CLM-011", "Phase 1 Milestone Payments: M-1 (40%), M-2 (40%), M-3 (20%).",
         "OFFICIAL_SOURCE", "Phase1_Scope_of_Work_Toor_Dal.docx: Table 45", "40/40/20 split on ₹1,73,000 base fee.",
         "M-1 = ₹69,200; M-2 = ₹69,200; M-3 = ₹34,600. Total = ₹1,73,000.", "Tied to inspectable deliverables, not cooking reduction guarantees.",
         "RETAIN_VERIFIED", "Proposal Sec 13, Commercial Basis Dossier")
    ]
    
    for row_idx, r in enumerate(claims_data, 5):
        is_zebra = (row_idx % 2 == 0)
        fill = ALT_FILL if is_zebra else None
        for col_idx, val in enumerate(r, 1):
            cell = ws1.cell(row=row_idx, column=col_idx, value=val)
            cell.font = DATA_FONT_BOLD if col_idx == 1 else DATA_FONT
            cell.border = THIN_BORDER
            if fill:
                cell.fill = fill
            if col_idx in [1, 3, 8]:
                cell.alignment = Alignment(horizontal="center", vertical="center")
            else:
                cell.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
        ws1.row_dimensions[row_idx].height = 24
        
    auto_fit_columns(ws1)
    
    # Sheet 2: Taxonomy Reference
    ws2 = wb.create_sheet(title="Epistemic_Taxonomy")
    ws2.views.sheetView[0].showGridLines = True
    ws2["A1"] = "FINNOAQ EPISTEMIC EVIDENCE TAXONOMY"
    ws2["A1"].font = TITLE_FONT
    
    tax_headers = ["Category Label", "Formal Definition", "Operational Rule in FinnoaQ Workspace"]
    for col_idx, h in enumerate(tax_headers, 1):
        cell = ws2.cell(row=3, column=col_idx, value=h)
        cell.fill = STEEL_FILL
        cell.font = WHITE_FONT_BOLD
        cell.border = THIN_BORDER
    
    tax_data = [
        ("MEASURED", "Physically measured with sample ID, verified standard method, date, and raw lab data.", "Highest certainty. Requires auditable instrument log."),
        ("SOURCE_REPORTED", "Reported by client or source document; not independently verified.", "Must be identified as preliminary client observation."),
        ("OFFICIAL_SOURCE", "Verified against an identifiable official statutory document (FSSAI, BIS, Codex).", "Mandates citation of exact regulation clause and year."),
        ("LITERATURE_SUPPORTED", "Supported by peer-reviewed literature with applicability limitations noted.", "Species and product format distinctions must be explicit."),
        ("CALCULATED", "Derived from explicitly identified inputs and reproducible first-principles math.", "Requires exact equations and zero rounding leakage."),
        ("SIMULATED", "Generated by a numerical model (Aspen, FEA), with convergence status stated.", "Cannot be represented as validated physical behavior."),
        ("ASSUMPTION", "Provisionally selected to perform an explicitly labelled scenario calculation.", "Must be subjected to sensitivity analysis."),
        ("HYPOTHESIS", "A possible causal explanation requiring empirical testing.", "Must compete with alternatives; null finding is valid."),
        ("NOT_VERIFIED", "Verification has not been completed.", "Flags missing data or unexecuted checks."),
        ("FAILED", "A test, check, or requirement did not pass.", "Immediate hold point; no silent continuation."),
        ("BLOCKED", "Execution prevented by external dependency, missing approval, or unavailable tool.", "Requires transparent stakeholder notification.")
    ]
    
    for row_idx, r in enumerate(tax_data, 4):
        for col_idx, val in enumerate(r, 1):
            cell = ws2.cell(row=row_idx, column=col_idx, value=val)
            cell.font = DATA_FONT_BOLD if col_idx == 1 else DATA_FONT
            cell.border = THIN_BORDER
            cell.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
        ws2.row_dimensions[row_idx].height = 22
    auto_fit_columns(ws2)
    
    wb.save(path)
    print(f"Generated: {path}")

# ==============================================================================
# 2. EXPERIMENT REGISTER WORKBOOK
# ==============================================================================
def create_experiment_register_wb(path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Experiment_Register"
    ws.views.sheetView[0].showGridLines = True
    
    ws["A1"] = "FINNOAQ TOOR DAL R&D — MASTER EXPERIMENT REGISTER"
    ws["A1"].font = TITLE_FONT
    ws["A2"] = "Governing Standard: Section 6 (Finding D) Replication Hierarchy | Status: PLANNED GATES"
    ws["A2"].font = SUBTITLE_FONT
    
    headers = [
        "Run ID", "Consignment Lot ID", "Material Origin & Variety",
        "Tier 2 Prep Batch", "Tier 3 Tech Repeat", "Treatment Code",
        "Reagent Description", "Dose (% w/w solids)", "Soak T (°C)", "Soak t (min)",
        "Primary Planned Measurement", "Run Status", "Raw Data Path", "Sign-Off"
    ]
    
    for col_idx, h in enumerate(headers, 1):
        cell = ws.cell(row=4, column=col_idx, value=h)
        cell.fill = NAVY_FILL
        cell.font = WHITE_FONT_BOLD
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = THIN_BORDER
    ws.row_dimensions[4].height = 28
    
    runs = []
    # Diagnostic Runs for Option 1: Lot TZ-01 (Tanzanian) & Lot IN-01 (Indian Benchmark)
    lots = [
        ("LOT-TZ-01", "Tanzanian Shinyanga Split (Target)"),
        ("LOT-IN-01", "Indian Latur Desi Split (Benchmark)")
    ]
    
    run_counter = 1
    for lot_id, lot_desc in lots:
        # Untreated Control (3 batches x 3 repeats)
        for b in ["B-01", "B-02", "B-03"]:
            for r in ["R-01", "R-02", "R-03"]:
                runs.append((
                    f"EXP-{run_counter:03d}", lot_id, lot_desc, b, r,
                    "TRT-00", "Untreated Control (As Received)", 0.0, 25.0, 0.0,
                    "Proximate + Raw Grain Hardness + Standard Cooking T90",
                    "PLANNED", "TBD_UPON_RECEIPT", "Sourish Senapati"
                ))
                run_counter += 1
                
        # Water-Only Control Soak (3 batches x 3 repeats)
        for b in ["B-01", "B-02", "B-03"]:
            for r in ["R-01", "R-02", "R-03"]:
                runs.append((
                    f"EXP-{run_counter:03d}", lot_id, lot_desc, b, r,
                    "TRT-01", "Pure Demineralized Water Soak", 0.0, 25.0, 30.0,
                    "Hydration Curve Mt + Texture Post-Dry + Cooking T90",
                    "PLANNED", "TBD_UPON_RECEIPT", "Sourish Senapati"
                ))
                run_counter += 1

    for row_idx, r in enumerate(runs, 5):
        is_zebra = (row_idx % 2 == 0)
        fill = ALT_FILL if is_zebra else None
        for col_idx, val in enumerate(r, 1):
            cell = ws.cell(row=row_idx, column=col_idx, value=val)
            cell.font = DATA_FONT_BOLD if col_idx in [1, 2, 6, 12] else DATA_FONT
            cell.border = THIN_BORDER
            if fill:
                cell.fill = fill
            if col_idx in [1, 2, 4, 5, 6, 8, 9, 10, 12]:
                cell.alignment = Alignment(horizontal="center", vertical="center")
            else:
                cell.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
        ws.row_dimensions[row_idx].height = 20
        
    auto_fit_columns(ws)
    wb.save(path)
    print(f"Generated: {path}")

# ==============================================================================
# 3. DELIVERABLE ACCEPTANCE CHECKLIST WORKBOOK
# ==============================================================================
def create_acceptance_checklist_wb(path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Acceptance_Checklist"
    ws.views.sheetView[0].showGridLines = True
    
    ws["A1"] = "FINNOAQ TOOR DAL R&D — DELIVERABLE ACCEPTANCE CHECKLIST"
    ws["A1"].font = TITLE_FONT
    ws["A2"] = "Verification Standard: Pre-Response Verification Gate & ASME/FSSAI Rules | Date: 20-Oct-2026"
    ws["A2"].font = SUBTITLE_FONT
    
    headers = [
        "Gate ID", "Deliverable Name & Scope", "Deliverable Type",
        "Governing Quality / Engineering Standard", "Verification Method & Evidence File",
        "Pass / Fail Threshold Criteria", "Actual Audit Result",
        "QA Status", "Lead Reviewer Sign-Off"
    ]
    
    for col_idx, h in enumerate(headers, 1):
        cell = ws.cell(row=4, column=col_idx, value=h)
        cell.fill = NAVY_FILL
        cell.font = WHITE_FONT_BOLD
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = THIN_BORDER
    ws.row_dimensions[4].height = 28
    
    gates = [
        ("GATE-01", "Phase 1 Diagnostic Proposal (Rev V10 DOCX/PDF)", "Client Deliverable",
         "SOW Concordance & Professional Visual QA", "Page-by-page PyMuPDF audit (audit_pdf_pages.py)",
         "14 pages; >=8pt font; zero empty pages; zero text clipping.", "14 pages fully populated, zero defects.",
         "PASSED", "Sourish Senapati"),
        ("GATE-02", "Mass Balance Closure & Scenario Reconciliation", "Engineering Calculation",
         "First-Principles Conservation of Mass", "MASS_AND_MOISTURE_BALANCE_AUDIT.md",
         "Mass in = Mass out (closure error <= 0.05%); 199.7 kg delta reconciled.", "0.000000% closure error across both scenarios.",
         "PASSED", "Sourish Senapati"),
        ("GATE-03", "10-Hypothesis Diagnostic Matrix (H1 to H10)", "Scientific Methodology",
         "Section 3 Directive (Criticism A)", "DIAGNOSTIC_HYPOTHESIS_AND_TEST_MATRIX.md",
         "Pectin crosslinking competes with 9 alternatives; null finding valid.", "All 10 hypotheses formalized with discriminating logic.",
         "PASSED", "Sourish Senapati"),
        ("GATE-04", "Three-Tier Replication Hierarchy", "Experimental Protocol",
         "Section 6 Directive (Criticism D)", "EXPERIMENTAL_DESIGN_AND_COOKING_TEST_PROTOCOL.md",
         "Separate source lots, prep batches, and technical repeats.", "Explicitly separated; 36 planned runs in register.",
         "PASSED", "Sourish Senapati"),
        ("GATE-05", "Chemical Treatment Specification & Controls", "Process Chemistry",
         "Section 5 Directive (Criticism C)", "EXPERIMENTAL_DESIGN_AND_COOKING_TEST_PROTOCOL.md",
         "Distinct formulations (NaHCO3, Citrate, Carbonate, Thermal) + Controls.", "Fully specified with grade, dosage, liquor ratio, wash.",
         "PASSED", "Sourish Senapati"),
        ("GATE-06", "Food Safety & Regulatory Compliance Matrix", "Regulatory Review",
         "FSSAI Food Additives & Pulses Standards", "LITERATURE_AND_REGULATORY_REGISTER.md",
         "FSSAI Table 3 GMP compliance; retained Na < 150 mg/100g.", "INS 500(ii) GMP confirmed; 53.9 mg Na/100g compliant.",
         "PASSED", "Sourish Senapati"),
        ("GATE-07", "Native SolidWorks CAD Models & Assemblies", "Mechanical Engineering",
         "ASME Y14.5 / Parametric Associativity", "SolidWorks COM verification (DR-101, TK-101, PLANT-15T)",
         "Genuine .SLDPRT / .SLDASM; no Python scripted raster mocks.", "SLDWORKS.exe connected; native parts and assembly verified.",
         "PASSED", "Sourish Senapati"),
        ("GATE-08", "Converged Aspen Plus Evaporator Simulation", "Process Simulation",
         "Thermodynamic Convergence (ELECNRTL)", "FINNOAQ_15TPD_CONVERGED.bkp",
         "Residuals < 1e-6; mass & energy closure; utility only.", "ELECNRTL converged; 1,000 kg/h feed; duty confirmed.",
         "PASSED", "Sourish Senapati"),
        ("GATE-09", "Persistent Workspace Governance Files", "Project Governance",
         "Section 14 Governance Directive", "AGENTS.md & .agents/rules/finnoaq_...",
         "Saved to disk and verified via read-back; links active.", "Files created, read-back verified, markdown lints resolved.",
         "PASSED", "Sourish Senapati"),
        ("GATE-10", "Phase 1 Scope Boundary & Fee Integrity", "Commercial Governance",
         "Authoritative SOW Table 40 & 45", "SCOPE_RECONCILIATION_AND_COMMERCIAL_BASIS.md",
         "Fixed fee ₹1,73,000; 40/40/20 milestones; factory design sequestered.", "Exact arithmetic match; factory design locked to Phase 4.",
         "PASSED", "Sourish Senapati")
    ]
    
    for row_idx, r in enumerate(gates, 5):
        is_zebra = (row_idx % 2 == 0)
        fill = ALT_FILL if is_zebra else None
        for col_idx, val in enumerate(r, 1):
            cell = ws.cell(row=row_idx, column=col_idx, value=val)
            cell.font = DATA_FONT_BOLD if col_idx in [1, 8] else DATA_FONT
            cell.border = THIN_BORDER
            if fill:
                cell.fill = fill
            if col_idx in [1, 3, 8]:
                cell.alignment = Alignment(horizontal="center", vertical="center")
            else:
                cell.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
        ws.row_dimensions[row_idx].height = 24
        
    auto_fit_columns(ws)
    wb.save(path)
    print(f"Generated: {path}")

if __name__ == "__main__":
    os.makedirs("docs/evidence", exist_ok=True)
    os.makedirs("docs/experiments", exist_ok=True)
    os.makedirs("docs/acceptance", exist_ok=True)
    
    create_evidence_matrix_wb(os.path.join("docs", "evidence", "EVIDENCE_MATRIX.xlsx"))
    create_experiment_register_wb(os.path.join("docs", "experiments", "EXPERIMENT_REGISTER.xlsx"))
    create_acceptance_checklist_wb(os.path.join("docs", "acceptance", "DELIVERABLE_ACCEPTANCE_CHECKLIST.xlsx"))

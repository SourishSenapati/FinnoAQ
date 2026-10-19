# -*- coding: utf-8 -*-
"""
V10 Proposal Content, Numerical and Regulatory Audit Test Suite
"""

import os
import re
import docx

def test_docx_audit():
    path = r"deliverables\PHASE1_TOOR_DAL_DIAGNOSTIC_RND_PROPOSAL_REV_V10.docx"
    assert os.path.exists(path), f"File not found: {path}"
    doc = docx.Document(path)
    
    # 1. Full text extraction
    full_text = []
    for p in doc.paragraphs:
        if p.text.strip():
            full_text.append(p.text.strip())
    for t in doc.tables:
        for r in t.rows:
            row_txt = [c.text.strip().replace('\n', ' ') for c in r.cells]
            full_text.append(' | '.join(row_txt))
    all_str = '\n'.join(full_text)
    
    print(f"Total extracted blocks: {len(full_text)}")
    print(f"Total characters: {len(all_str)}")
    
    # 2. Check prohibited strings (Anti-Fraud / Quality Governance)
    prohibited = [
        "Harakuntir", # Must be Harakutir (no 'n')
        "five-sigma",
        "proven mechanism",
        "guaranteed cooking",
        "ready for construction",
        "zero error"
    ]
    for p_str in prohibited:
        matches = re.findall(re.escape(p_str), all_str, re.IGNORECASE)
        print(f"Check '{p_str}': {len(matches)} occurrences")
        assert len(matches) == 0, f"Prohibited string '{p_str}' found in document!"

    # 3. Check Mandatory Identities
    assert "Harakutir" in all_str, "Client Harakutir missing"
    assert "FinnoaQ" in all_str, "FinnoaQ missing"
    assert "Sourish Senapati" in all_str, "Lead Engineer missing"
    assert "Dr. Archisman Ray" in all_str, "Academic Guidance missing"
    assert "IIT Kharagpur" in all_str, "IIT Kharagpur missing"
    
    # 4. Check Mandatory Commercial Numbers
    assert "1,73,000" in all_str, "Fee 1,73,000 missing"
    assert "69,200" in all_str, "Milestone 69,200 missing"
    assert "34,600" in all_str, "Milestone 34,600 missing"
    assert "2,00,000" in all_str, "Ceiling 2,00,000 missing"
    assert "20–25" in all_str or "20-25" in all_str, "Timeline 20-25 days missing"
    
    # 5. Check Mass Balance Numbers
    assert "13,425" in all_str, "Dry solids 13,425 missing"
    assert "13,241" in all_str, "Dry solids 13,241 missing"
    assert "14,712.222" in all_str or "14,712.22" in all_str, "Product mass 14,712.222 missing"
    assert "18,390.278" in all_str or "18,390.28" in all_str, "Scenario B1 pre-dryer missing"
    assert "3,678.056" in all_str or "3,678.06" in all_str, "Scenario B1 evap missing"
    assert "18,590" in all_str, "Scenario B2 pre-dryer missing"
    assert "28.77" in all_str, "Scenario B2 implied moisture missing"
    assert "3,877.778" in all_str or "3,877.78" in all_str, "Scenario B2 evap missing"
    assert "1.3706%" in all_str or "1.370577%" in all_str or "1.37%" in all_str, "Solids loss percentage missing"
    
    # 6. Check Chemical Reagent Segregation
    assert "NaHCO₃" in all_str or "NaHCO3" in all_str, "Bicarbonate missing"
    assert "Na₂CO₃" in all_str or "Na2CO3" in all_str, "Carbonate missing"
    assert "Na₃C₆H₅O₇" in all_str or "Na3C6H5O7" in all_str or "Trisodium Citrate" in all_str, "Citrate missing"
    assert "Hydrothermal" in all_str, "Hydrothermal missing"
    assert "PHASE 2 CANDIDATE — NOT TESTED" in all_str, "Candidate phase tag missing"

    # 7. Check Sample Allocation & Replicate Plan
    assert "2,550 g" in all_str or "2.55 kg" in all_str, "Sample mass budget missing"
    assert "Option 1" in all_str and "Option 2" in all_str, "Options 1 and 2 missing"
    
    # 8. Check Food Safety & Regulatory Limits
    assert "2.4.6.15" in all_str, "FSSAI pulse standard clause missing"
    assert "14.0%" in all_str, "FSSAI moisture limit missing"
    assert "150 mg" in all_str, "Sodium target missing"
    
    print("\nALL CONTENT AND NUMERICAL CHECKS PASSED PERFECTLY!")

if __name__ == "__main__":
    test_docx_audit()

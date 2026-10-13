import sys
import os
import docx

def dump_docx(filename, out_txt):
    print(f"Reading {filename}...")
    doc = docx.Document(filename)
    with open(out_txt, "w", encoding="utf-8") as f:
        f.write(f"# DUMP OF {filename}\n\n")
        f.write("## PARAGRAPHS\n\n")
        for i, p in enumerate(doc.paragraphs):
            t = p.text.strip()
            if t:
                f.write(f"[{i}] {t}\n")
        f.write("\n## TABLES\n\n")
        for t_idx, table in enumerate(doc.tables):
            f.write(f"### Table {t_idx+1}\n")
            for r_idx, row in enumerate(table.rows):
                row_vals = [cell.text.strip().replace('\n', ' ') for cell in row.cells]
                f.write(f"| " + " | ".join(row_vals) + " |\n")
            f.write("\n")
    print(f"Wrote dump to {out_txt}")

if __name__ == "__main__":
    dump_docx("Phase1_Scope_of_Work_Toor_Dal.docx", "phase1_sow_dump.txt")
    dump_docx("Toor_Dal_Quick_Cooking_RD_Master_Technical_Scope.docx", "quick_cooking_master_scope_dump.txt")

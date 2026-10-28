# -*- coding: utf-8 -*-
"""
V10 Master Proposal Generator: Phase 1 Comparative Diagnostic R&D Study
Client: Client Organization (Kolkata) | Lead Engineer: Sourish Senapati (FinnoaQ)
Academic Guidance: Technical Advisory Panel
"""

import os
import sys
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import qn, nsdecls

# Color Palette Constants (Professional Engineering Consultancy)
HEX_NAVY = "1A365D"       # Primary Dark Navy
HEX_BLUE = "2B6CB0"       # Secondary Steel Blue
HEX_DARK_GREY = "2D3748"  # Body Text Dark Charcoal
HEX_LIGHT_GREY = "F7FAFC" # Table Alternate Fill
HEX_BORDER_GREY = "CBD5E0"# Subtle Border Grey
HEX_CALLOUT_BG = "EDF2F7" # Callout Box Fill
HEX_ALERT_BORDER = "3182CE"# Callout Left Border
HEX_WHITE = "FFFFFF"

COLOR_NAVY = RGBColor(26, 54, 93)
COLOR_BLUE = RGBColor(43, 108, 176)
COLOR_DARK_GREY = RGBColor(45, 55, 72)
COLOR_MUTED = RGBColor(113, 128, 150)

def set_cell_margins(cell, top=120, bottom=120, left=150, right=150):
    """Set padding for a single table cell (twips: 20 twips = 1 pt)."""
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def set_cell_shading(cell, color_hex):
    """Set background color of a cell."""
    shading_elm = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{color_hex}"/>')
    cell._tc.get_or_add_tcPr().append(shading_elm)

def set_table_borders(table, color="CBD5E0", sz="4", val="single"):
    """Set subtle, professional borders for the entire table."""
    tblPr = table._tbl.tblPr
    borders = parse_xml(
        f'<w:tblBorders {nsdecls("w")}>'
        f'  <w:top w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'  <w:bottom w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'  <w:left w:val="none"/>'
        f'  <w:right w:val="none"/>'
        f'  <w:insideH w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'  <w:insideV w:val="none"/>'
        f'</w:tblBorders>'
    )
    tblPr.append(borders)

def make_row_header(row, bg_hex=HEX_NAVY):
    """Style a table row as a repeating header row."""
    trPr = row._tr.get_or_add_trPr()
    trPr.append(parse_xml(f'<w:tblHeader {nsdecls("w")}/>'))
    trPr.append(parse_xml(f'<w:cantSplit {nsdecls("w")}/>'))
    for cell in row.cells:
        set_cell_shading(cell, bg_hex)
        set_cell_margins(cell, top=140, bottom=140, left=150, right=150)
        for p in cell.paragraphs:
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            for run in p.runs:
                run.font.bold = True
                run.font.color.rgb = RGBColor(255, 255, 255)
                run.font.size = Pt(9.5)
                run.font.name = "Calibri"

def style_data_row(row, is_zebra=False, font_size=9.0):
    """Style data row with padding, cantSplit, and optional zebra background."""
    trPr = row._tr.get_or_add_trPr()
    trPr.append(parse_xml(f'<w:cantSplit {nsdecls("w")}/>'))
    bg = HEX_LIGHT_GREY if is_zebra else HEX_WHITE
    for cell in row.cells:
        if is_zebra:
            set_cell_shading(cell, bg)
        set_cell_margins(cell, top=100, bottom=100, left=150, right=150)
        for p in cell.paragraphs:
            for run in p.runs:
                run.font.size = Pt(font_size)
                run.font.name = "Calibri"
                run.font.color.rgb = COLOR_DARK_GREY

def add_callout(doc, text, heading="SCOPE SPECIFICATION & DIAGNOSTIC BOUNDARY", border_color=HEX_ALERT_BORDER, bg_color=HEX_CALLOUT_BG):
    """Insert a highlighted note/callout box."""
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = False
    tbl.columns[0].width = Inches(7.0)
    
    cell = tbl.cell(0, 0)
    set_cell_shading(cell, bg_color)
    set_cell_margins(cell, top=140, bottom=140, left=200, right=160)
    
    # Left thick border, no other borders
    tcPr = cell._tc.get_or_add_tcPr()
    borders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>'
        f'  <w:left w:val="single" w:sz="24" w:space="0" w:color="{border_color}"/>'
        f'  <w:top w:val="none"/>'
        f'  <w:right w:val="none"/>'
        f'  <w:bottom w:val="none"/>'
        f'</w:tcBorders>'
    )
    tcPr.append(borders)
    
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(4)
    run_h = p.add_run(f"[{heading}]\n")
    run_h.font.bold = True
    run_h.font.size = Pt(9.5)
    run_h.font.color.rgb = COLOR_NAVY
    run_h.font.name = "Calibri"
    
    run_b = p.add_run(text)
    run_b.font.size = Pt(9.0)
    run_b.font.italic = True
    run_b.font.color.rgb = COLOR_DARK_GREY
    run_b.font.name = "Calibri"
    
    doc.add_paragraph().paragraph_format.space_after = Pt(4)

print("Helper functions compiled successfully.")

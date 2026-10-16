"""
Build V8 Client Project Proposal (DOCX & PDF)
Document Ref: FNQ-PR-2026-001 Rev 8.0
Lead Chemical & Process Engineer: Sourish Senapati (FinnoaQ Agro-Industrial Technologies Ltd.)
Academic Guidance: Dr. Archisman Ray (IIT Kharagpur)
"""

import os
import sys
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import qn, nsdecls
import win32com.client

BASE_DIR = os.path.abspath(".")
DELIV_DIR = os.path.join(BASE_DIR, "deliverables", "V8_CLIENT_ONBOARDING_AND_ENGINEERING")
os.makedirs(DELIV_DIR, exist_ok=True)

DOCX_OUT = os.path.join(DELIV_DIR, "CLIENT_PROJECT_PROPOSAL.docx")
PDF_OUT = os.path.join(DELIV_DIR, "CLIENT_PROJECT_PROPOSAL.pdf")

print(f"Generating V8 Client Proposal: {DOCX_OUT}", flush=True)

doc = docx.Document()

# Page Setup: Standard A4 (8.27 x 11.69 inches) with 0.8 in margins
for sec in doc.sections:
    sec.page_width = Inches(8.27)
    sec.page_height = Inches(11.69)
    sec.top_margin = Inches(0.8)
    sec.bottom_margin = Inches(0.8)
    sec.left_margin = Inches(0.8)
    sec.right_margin = Inches(0.8)

# Color Palette: Executive Engineering Corporate
CLR_NAVY = RGBColor(15, 23, 42)       # Slate 900
CLR_BLUE = RGBColor(30, 58, 138)      # Blue 900
CLR_DARK = RGBColor(30, 41, 59)       # Slate 800
CLR_MUTED = RGBColor(71, 85, 105)     # Slate 600
HEX_BORDER = "CBD5E1"                 # Slate 300
HEX_HDR_BG = "F1F5F9"                 # Slate 100
HEX_ACCENT_BG = "EFF6FF"              # Blue 50

# Base Styles
style_normal = doc.styles['Normal']
style_normal.font.name = 'Calibri'
style_normal.font.size = Pt(10)
style_normal.font.color.rgb = CLR_DARK

def format_para(p, space_before=0, space_after=4, line_spacing=1.15):
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = line_spacing

def add_header(title, level=1):
    p = doc.add_paragraph()
    run = p.add_run(title)
    run.bold = True
    if level == 1:
        run.font.size = Pt(14)
        run.font.color.rgb = CLR_NAVY
        format_para(p, space_before=12, space_after=4)
        # Add a subtle bottom border or underline via XML
        pPr = p._p.get_or_add_pPr()
        pBdr = parse_xml(f'<w:pBdr {nsdecls("w")}><w:bottom w:val="single" w:sz="6" w:space="2" w:color="3B82F6"/></w:pBdr>')
        pPr.append(pBdr)
    elif level == 2:
        run.font.size = Pt(11.5)
        run.font.color.rgb = CLR_BLUE
        format_para(p, space_before=8, space_after=3)
    elif level == 3:
        run.font.size = Pt(10.5)
        run.font.color.rgb = CLR_DARK
        format_para(p, space_before=6, space_after=2)
    return p

def set_cell_properties(cell, bg_hex=None, top_pad=40, bot_pad=40, l_pad=60, r_pad=60):
    tcPr = cell._tc.get_or_add_tcPr()
    if bg_hex:
        shd = parse_xml(f'<w:shd {nsdecls("w")} w:val="clear" w:color="auto" w:fill="{bg_hex}"/>')
        tcPr.append(shd)
    tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top_pad}" w:type="dxa"/><w:bottom w:w="{bot_pad}" w:type="dxa"/><w:left w:w="{l_pad}" w:type="dxa"/><w:right w:w="{r_pad}" w:type="dxa"/></w:tcMar>')
    tcPr.append(tcMar)

def style_table(table, col_widths, col_alignments):
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    tblPr = table._tbl.tblPr
    borders = parse_xml(
        f'<w:tblBorders {nsdecls("w")}>'
        f'  <w:top w:val="single" w:sz="6" w:space="0" w:color="{HEX_BORDER}"/>'
        f'  <w:bottom w:val="single" w:sz="6" w:space="0" w:color="{HEX_BORDER}"/>'
        f'  <w:left w:val="none"/>'
        f'  <w:right w:val="none"/>'
        f'  <w:insideH w:val="single" w:sz="4" w:space="0" w:color="{HEX_BORDER}"/>'
        f'  <w:insideV w:val="none"/>'
        f'</w:tblBorders>'
    )
    tblPr.append(borders)
    for r_idx, row in enumerate(table.rows):
        trPr = row._tr.get_or_add_trPr()
        trPr.append(parse_xml(f'<w:cantSplit {nsdecls("w")}/>'))
        if r_idx == 0:
            trPr.append(parse_xml(f'<w:tblHeader {nsdecls("w")}/>'))
        for c_idx, cell in enumerate(row.cells):
            cell.width = Inches(col_widths[c_idx])
            bg = HEX_HDR_BG if r_idx == 0 else (HEX_ACCENT_BG if r_idx % 2 == 1 else "FFFFFF")
            set_cell_properties(cell, bg_hex=bg)
            for p in cell.paragraphs:
                p.alignment = col_alignments[c_idx]
                p.paragraph_format.space_before = Pt(2)
                p.paragraph_format.space_after = Pt(2)
                p.paragraph_format.line_spacing = 1.1
                for run in p.runs:
                    run.font.size = Pt(8.5 if r_idx > 0 else 9.0)
                    if r_idx == 0:
                        run.bold = True
                        run.font.color.rgb = CLR_NAVY

# ==============================================================================
# TITLE & METADATA BLOCK
# ==============================================================================
p_meta = doc.add_paragraph()
r_meta = p_meta.add_run("FINNOAQ AGRO-INDUSTRIAL TECHNOLOGIES LTD. | COMMERCIAL PROPOSAL & SOW")
r_meta.bold = True
r_meta.font.size = Pt(8.5)
r_meta.font.color.rgb = CLR_BLUE
format_para(p_meta, space_before=0, space_after=2)

p_title = doc.add_paragraph()
r_title = p_title.add_run("Comparative R&D Investigation & Cooking-Performance Optimization for Imported Tanzanian Toor Dal (Cajanus cajan)")
r_title.bold = True
r_title.font.size = Pt(17)
r_title.font.color.rgb = CLR_NAVY
format_para(p_title, space_before=2, space_after=4)

p_sub = doc.add_paragraph()
r_sub = p_sub.add_run("Phase 1 Diagnostic Research Scope of Work, Commercial Terms & Technical Framework | Ref: FNQ-PR-2026-001 (Rev 8.0)")
r_sub.font.size = Pt(10)
r_sub.font.color.rgb = CLR_MUTED
format_para(p_sub, space_before=0, space_after=10)

# Project Metadata Table
meta_tbl = doc.add_table(rows=4, cols=2)
meta_data = [
    ("Client / Sponsor:", "Commercial Pulse Processing Facility Management & Technical Team (India)"),
    ("Lead Chemical & Process Engineer:", "Sourish Senapati, FinnoaQ Agro-Industrial Technologies Ltd. (Kolkata, WB)"),
    ("Academic Guidance & Scientific Advisory:", "Dr. Archisman Ray, Senior Research Fellow, Dept. of Chemical Engineering, IIT Kharagpur"),
    ("Project Date & Execution Baseline:", "16 October 2026 | Execution Mode: Offline-First Native Engineering | Revision 8.0")
]
for i, (k, v) in enumerate(meta_data):
    r = meta_tbl.rows[i]
    r.cells[0].paragraphs[0].add_run(k).bold = True
    r.cells[1].paragraphs[0].add_run(v)
style_table(meta_tbl, [2.2, 4.47], [WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.LEFT])

p_sep = doc.add_paragraph()
format_para(p_sep, space_before=6, space_after=6)

# ==============================================================================
# 1. EXECUTIVE SUMMARY & COMMERCIAL CONTEXT
# ==============================================================================
add_header("1. Executive Summary & Commercial Rationale")

p1 = doc.add_paragraph(
    "Imported Tanzanian split, dehulled pigeon pea (Cajanus cajan)—commonly traded as Toor Dal—presents a compelling raw material "
    "procurement price advantage of approximately ₹18–22/kg relative to domestic Indian benchmark varieties. However, commercial uptake "
    "among Indian industrial millers and retail consumers is severely constrained by an acute cooking latency defect: under standard household "
    "pressure-cooking conditions, Tanzanian toor dal requires approximately 6 whistles (~45–55 minutes total cooking cycle) compared to 3 whistles "
    "(~20–25 minutes) for comparable domestic Indian toor dal. Even after extended thermal boiling, the imported grains frequently retain a persistent "
    "hard core, exhibit non-uniform softening, and display poor culinary mouthfeel."
)
format_para(p1, space_before=3, space_after=4)

p2 = doc.add_paragraph(
    "This proposal establishes an authoritative, five-stage Phase 1 Diagnostic Research Program designed to scientifically determine the root "
    "biophysical, structural, and chemical mechanisms responsible for this cooking-time disparity. Operating under strict engineering governance, "
    "this study will generate empirical laboratory evidence across proximate composition, cell-wall pectin fractionation, mineral cross-linking, "
    "starch gelatinization, and hydration kinetics to provide a defensible foundation for subsequent commercial conditioning and industrial processing."
)
format_para(p2, space_before=0, space_after=8)

# ==============================================================================
# 2. PHASE 1 DETAILED SCOPE OF WORK (11 ANALYTICAL MODULES)
# ==============================================================================
add_header("2. Phase 1 Technical Investigation Scope of Work")

p3 = doc.add_paragraph(
    "The Phase 1 investigation replaces qualitative household observations with rigorous, objective, and reproducible analytical methodologies "
    "adhering to AOAC, ISO, BIS, and FSSAI standards. Both the imported Tanzanian test lot and a representative domestic Indian benchmark lot "
    "will be evaluated simultaneously across 11 synchronized analytical work packages:"
)
format_para(p3, space_before=2, space_after=6)

sow_tbl = doc.add_table(rows=12, cols=4)
sow_headers = ["Module Ref", "Work Package Description", "Analytical Methodology & Instrumentation", "Indicative Fee (INR)"]
for c_idx, h in enumerate(sow_headers):
    sow_tbl.rows[0].cells[c_idx].paragraphs[0].add_run(h)

sow_rows = [
    ("WP-01", "Sample Preparation & Physical Characterization", "Grain dimensions (L/W ratio), 1000-grain weight, bulk density, broken %, defect count, raw hardness via digital micrometer and sieving.", "₹20,000"),
    ("WP-02", "Proximate Chemical Composition", "Standard AOAC / FSSAI protocols: moisture (vacuum oven), crude protein (Kjeldahl), crude fat (Soxhlet), crude fiber, total ash, and available carbohydrate.", "₹7,000"),
    ("WP-03", "Starch & Amylose Profiling", "Total starch determination, amylose-to-amylopectin ratio via iodine spectrophotometry, and enzymatic starch availability assays.", "₹5,000"),
    ("WP-04", "DSC Gelatinization & Thermal Behavior", "Differential Scanning Calorimetry (DSC): onset (To), peak (Tp), concluding (Tc) gelatinization temperatures, transition enthalpy (ΔH), and pasting behavior.", "₹20,000"),
    ("WP-05", "Hydration Kinetics & Water Absorption Index", "Multi-point isothermal soaking curves (0, 15, 30, 60, 120, 240, 480 min) at 25°C, 45°C, and 65°C; gravimetric Water Absorption Index (WAI) and volume swelling.", "₹10,000"),
    ("WP-06", "Pectin & Cell-Wall Fractionation", "Sequential chemical fractionation: water-soluble pectin, chelator-soluble pectin (EDTA), and alkali-soluble pectin; degree of esterification; cell-wall rigidity analysis.", "₹18,000"),
    ("WP-07", "Mineral Profile & Phytate Interactions", "ICP-OES / AAS determination of divalent cations (Ca²⁺, Mg²⁺), iron, phosphorus, and spectrophotometric phytic acid quantification; phytate:cation molar ratios.", "₹16,000"),
    ("WP-08", "Instrumental Hardness & Texture Profile", "Texture Profile Analysis (TPA) on a calibrated Lloyd/Stable Micro Systems Texture Analyzer: raw breaking force, soaked firmness, cooked peak penetration resistance.", "₹17,000"),
    ("WP-09", "Standardized Cooking-Performance Trials", "Instrumented laboratory pressure vessel trials with calibrated temperature/pressure logging; determination of objective cooking endpoint (cooked core % < 2.0%).", "₹20,000"),
    ("WP-10", "Integrated Data Analysis & Root-Cause Matrix", "Multi-factor statistical correlation and analysis of variance (ANOVA) linking physicochemical parameters to cooking latency; candidate ranking.", "₹20,000"),
    ("WP-11", "Comprehensive Technical R&D Report", "Publication-grade final dossier synthesizing all raw data, statistical matrices, biophysical mechanisms, and formal Phase 2 processing recommendations.", "₹20,000"),
]

for r_idx, row_data in enumerate(sow_rows):
    row = sow_tbl.rows[r_idx + 1]
    for c_idx, val in enumerate(row_data):
        row.cells[c_idx].paragraphs[0].add_run(val)

style_table(sow_tbl, [0.8, 2.3, 2.77, 0.8], [WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.RIGHT])

# ==============================================================================
# 3. SCIENTIFIC EVIDENCE HIERARCHY & BOUNDARIES
# ==============================================================================
add_header("3. Scientific Evidence Hierarchy & Interpretation Principles")

p4 = doc.add_paragraph(
    "To ensure scientific and commercial defensibility, FinnoaQ enforces a strict four-tier evidence hierarchy. No causal mechanism will "
    "be reported as established fact unless it is experimentally demonstrated through replicated wet-lab testing:"
)
format_para(p4, space_before=2, space_after=4)

eh_tbl = doc.add_table(rows=5, cols=3)
eh_headers = ["Evidence Level", "Operational Definition", "Project Governance Rule"]
for c_idx, h in enumerate(eh_headers):
    eh_tbl.rows[0].cells[c_idx].paragraphs[0].add_run(h)

eh_rows = [
    ("Level 1: Measured Result", "Direct empirical physical or chemical measurement obtained from verified laboratory instruments under standard operating procedures.", "Reported with raw data, standard deviation, and analytical method citation."),
    ("Level 2: Observed Association", "Statistically significant correlation observed between a measured parameter (e.g., pectin-associated Ca²⁺) and cooking time.", "Categorized as 'Strongly Associated Factor'; cannot assert standalone causality."),
    ("Level 3: Scientific Hypothesis", "Biophysical or chemical mechanism proposed to explain observed associations (e.g., middle lamella pectate cross-linking).", "Explicitly labeled 'HYPOTHESIS'; carried forward for Phase 2 validation."),
    ("Level 4: Phase 2 Validation", "Controlled intervention trial demonstrating that targeted treatment reverses the defect under industrial conditions.", "Required before commercial process guarantees or scale-up commitments.")
]
for r_idx, row_data in enumerate(eh_rows):
    row = eh_tbl.rows[r_idx + 1]
    for c_idx, val in enumerate(row_data):
        row.cells[c_idx].paragraphs[0].add_run(val)
style_table(eh_tbl, [1.5, 3.2, 1.97], [WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.LEFT])

p5 = doc.add_paragraph(
    "Important Scope Boundary: Phase 1 is strictly a diagnostic R&D investigation. It does not provide commercial product guarantees, "
    "commercial processing parameters, or final construction engineering. These activities are addressed in subsequent phased offerings."
)
p5.runs[0].italic = True
format_para(p5, space_before=4, space_after=8)

# ==============================================================================
# 4. CANDIDATE TECHNOLOGIES FOR PHASE 2 SCREENING
# ==============================================================================
add_header("4. Candidate Interventions for Follow-On Phase 2 Evaluation")

p6 = doc.add_paragraph(
    "Based on Phase 1 diagnostic findings, Phase 2 will experimentally screen targeted interventions to accelerate cooking times. "
    "The relevance of each candidate will be determined strictly by the empirical Phase 1 evidence:"
)
format_para(p6, space_before=2, space_after=4)

cand_tbl = doc.add_table(rows=5, cols=3)
cand_headers = ["Intervention Category", "Target Mechanism & Physiological Action", "Regulatory & Processing Considerations"]
for c_idx, h in enumerate(cand_headers):
    cand_tbl.rows[0].cells[c_idx].paragraphs[0].add_run(h)

cand_rows = [
    ("Food-Grade Mineral Conditioning", "Alkaline salt soaking (Sodium Bicarbonate, Sodium Citrate) to chelate bound Ca²⁺/Mg²⁺ from insoluble pectates in the middle lamella, promoting cotyledon cell separation.", "Strict compliance with FSSAI Food Safety and Standards Regulations (Table 3 GMP limits); leaves zero chemical off-flavor or residue."),
    ("Controlled Hydrothermal Tempering", "Precision temperature-programmed soaking (45°C–65°C) to accelerate moisture diffusion without inducing premature starch gelatinization or cotyledon cracking.", "Requires continuous closed-loop water temperature control; minimal solids leaching (<1.5% w/w)."),
    ("Targeted Thermal / Microwave Pre-Treatment", "Rapid dry dielectric or infrared heating to create micro-fissures in dense cotyledon cell layers, facilitating rapid capillary water ingress.", "Requires precise exposure limits to prevent protein denaturation, darkening, or seed coat discoloration."),
    ("Enzymatic Cell-Wall Attenuation", "Pectinase / cellulase micro-dosing to selectively hydrolyze structural cell-wall polysaccharides and enhance cellular permeability.", "Enzyme inactivation requirements; economics and processing costs evaluated in techno-economic study.")
]
for r_idx, row_data in enumerate(cand_rows):
    row = cand_tbl.rows[r_idx + 1]
    for c_idx, val in enumerate(row_data):
        row.cells[c_idx].paragraphs[0].add_run(val)
style_table(cand_tbl, [1.8, 2.87, 2.0], [WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.LEFT])

p_sep2 = doc.add_paragraph()
format_para(p_sep2, space_before=4, space_after=4)

# ==============================================================================
# 5. COMMERCIAL TERMS, MILESTONES & SCHEDULE
# ==============================================================================
add_header("5. Commercial Terms, Turnaround Schedule & Milestones")

p7 = doc.add_paragraph(
    "FinnoaQ operates on an auditable milestone-based commercial structure. The total professional fee for the complete Phase 1 "
    "Diagnostic R&D Program is ₹1,73,000 (Rupees One Lakh Seventy-Three Thousand Only) plus applicable GST:"
)
format_para(p7, space_before=2, space_after=6)

comm_tbl = doc.add_table(rows=4, cols=4)
comm_headers = ["Project Milestone", "Scope Description & Deliverable", "Payment %", "Fee Amount (INR)"]
for c_idx, h in enumerate(comm_headers):
    comm_tbl.rows[0].cells[c_idx].paragraphs[0].add_run(h)

comm_rows = [
    ("Milestone 1: Project Initiation", "Formal execution of onboarding agreement; assignment of unique lot custody tracking numbers; physical sample receipt and preparation.", "40%", "₹69,200"),
    ("Milestone 2: Laboratory Completion", "Completion of all physical, proximate, starch, thermal (DSC), hydration, pectin, mineral (ICP-OES), texture, and standardized cooking trials; interim data transmittal.", "40%", "₹69,200"),
    ("Milestone 3: Final Report Release", "Delivery and formal presentation of Comprehensive Technical R&D Report, Statistical Root-Cause Correlation Matrix, and Prioritized Phase 2 Technology Recommendations.", "20%", "₹34,600"),
]
for r_idx, row_data in enumerate(comm_rows):
    row = comm_tbl.rows[r_idx + 1]
    for c_idx, val in enumerate(row_data):
        row.cells[c_idx].paragraphs[0].add_run(val)
style_table(comm_tbl, [1.6, 3.27, 0.8, 1.0], [WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.RIGHT])

p8 = doc.add_paragraph(
    "Key Commercial Conditions:\n"
    "• Commercial Working Ceiling: For procurement planning, an expenditure ceiling of up to ₹2,00,000 + applicable GST is established. Any additional specialized testing requires prior written client authorization.\n"
    "• Taxes: Goods and Services Tax (GST) will be charged additionally as applicable under prevailing statutory rates.\n"
    "• Turnaround Time: 20 to 25 business days from verified receipt of representative samples at our Kolkata analytical reception.\n"
    "• Fast-Track Interim Summary: An executive Stage 1 preliminary data brief will be delivered within 7 business days of sample arrival.\n"
    "• Offer Validity: This commercial proposal is valid for 30 calendar days from the date of issuance."
)
format_para(p8, space_before=4, space_after=8)

# ==============================================================================
# 6. SAMPLE DISPATCH & CHAIN OF CUSTODY
# ==============================================================================
add_header("6. Sample Dispatch & Chain of Custody Protocol")

p9 = doc.add_paragraph(
    "To ensure experimental validity and statistical rigor, the client will supply representative sample lots adhering to the following protocol:\n"
    "1. Required Quantity: 2.5 kg to 3.0 kg of the target imported Tanzanian commercial toor dal lot, and 2.5 kg to 3.0 kg of standard domestic Indian benchmark dal.\n"
    "2. Packaging: Samples must be double-bagged in food-grade, moisture-barrier polyethylene bags and sealed in rigid outer cartons.\n"
    "3. Lot Documentation: Each consignment must include batch/lot ID, port of entry, date of import/milling, and known storage conditions.\n"
    "4. Consignment Address: Central Laboratory Reception, FinnoaQ Agro-Industrial Technologies Ltd., Kolkata, West Bengal (Attn: Sourish Senapati).\n"
    "5. Retention: A 500 g sealed reference sample from each lot will be securely archived at 20°C / 55% RH for 12 months."
)
format_para(p9, space_before=2, space_after=8)

# ==============================================================================
# 7. PHASED ROADMAP: PHASE 2A & 15 TPD PLANT CONCEPT
# ==============================================================================
add_header("7. Industrial Scale-Up Roadmap: Phased Integration")

p10 = doc.add_paragraph(
    "To prevent premature capital expenditure, FinnoaQ structures industrial pulse processing into distinct, de-risked decision gates:\n"
    "• Phase 1 (Current Base Scope): Diagnostic baseline & biophysical root-cause identification (₹1.73 Lakh + GST, 20–25 days).\n"
    "• Optional Phase 2A (Laboratory Validation): Controlled formulation screening across candidate conditioning agents; bench-scale cooking confirmation; sensory, nutritional, and shelf-life assessment.\n"
    "• Optional Phase 2B (Commercial Plant FEED): Conceptual Front-End Engineering Design for a continuous 15,000 kg/day (15 TPD) industrial dal conditioning plant. FinnoaQ maintains an integrated 3D SolidWorks parametric model (PLANT-15T, 12 interconnected systems, 27.08 tonnes B-rep assembly), continuous fluid-bed dryer psychrometric models, and converged Aspen Plus V14 Zero-Liquid-Discharge (ZLD) falling-film evaporator balances (steam economy 2.38). This engineering baseline will be tailored to the client's verified site layout upon successful completion of Phase 1."
)
format_para(p10, space_before=2, space_after=6)

# ==============================================================================
# 8. MASTER ENGINEERING DRAWING ANNEXURES (15 TPD PLANT CONCEPT)
# ==============================================================================
add_header("8. Master Engineering Drawing Annexures (15 TPD Industrial Plant Concept)")

p_anx = doc.add_paragraph(
    "To substantiate the industrial viability of the proposed continuous hydration and drying technology, FinnoaQ maintains "
    "a fully developed Front-End Engineering Design (FEED) package. The following authentic engineering drawings, generated via "
    "Autodesk AutoCAD 2025 (accoreconsole native vector engine) and SolidWorks 2025 SP01.2 parametric solid modeling, are attached "
    "as official project annexures:"
)
format_para(p_anx, space_before=2, space_after=4)

anx_tbl = doc.add_table(rows=11, cols=4)
anx_headers = ["Annexure Ref", "Drawing Number & Title", "Design Standard & Scope Description", "Format & Native Tool Provenance"]
for c_idx, h in enumerate(anx_headers):
    anx_tbl.rows[0].cells[c_idx].paragraphs[0].add_run(h)

anx_rows = [
    ("Annexure A", "DWG-FIN-15T-001\nProcess Flow Diagram (PFD)", "12 reconciled streams, material balance flags, equipment silhouettes, CIP & ZLD connections.", "AutoCAD 2025 (Vector PDF)\nISO A2 (monochrome.ctb)"),
    ("Annexure B", "DWG-FIN-15T-002\nHeat & Material Balance Schedule", "Full A2 stream table (solids, water, enthalpy), exact 0.000% dry solids closure, utility duties.", "AutoCAD 2025 (Vector PDF)\nISO A2 (monochrome.ctb)"),
    ("Annexure C", "DWG-FIN-15T-003\nP&ID: Soaking Battery & Dosing Skid", "ISA-5.1 dual 7.5m3 soak tanks TK-101A/B, dosing skid SK-100, LIC/TIC/AIC/ZIC control loops.", "AutoCAD 2025 (Vector PDF)\nISO A2 (monochrome.ctb)"),
    ("Annexure D", "DWG-FIN-15T-004\nSite Master Plot Plan & Hygienic Zoning", "75m x 45m plot, 52m x 24m building, 4 ISO 22000 hygiene zones, room HVAC & area schedule.", "AutoCAD 2025 (Vector PDF)\nISO A2 (monochrome.ctb)"),
    ("Annexure E", "DWG-FIN-15T-007\nEquipment GA: Soak Vessel TK-101A/B", "ASME Section VIII Div 1, 2000mm ID, 60° cone, 4.0 kW hydrofoil agitator, nozzle schedule N1-N6/M1.", "AutoCAD 2025 + SolidWorks\nISO A2 (monochrome.ctb)"),
    ("Annexure F", "DWG-FIN-15T-008\nEquipment GA: Fluid Bed Dryer DR-101", "14m x 1m SS304 deck, 3 psychrometric zones, dual NFPA 68 burst panels, 8 vibration springs.", "AutoCAD 2025 (Vector PDF)\nISO A2 (monochrome.ctb)"),
    ("Annexure G", "DWG-FIN-15T-010\nP&ID: Triple-Effect ZLD Evaporator", "25 KLD falling-film calandrias, vacuum condenser, distillate recovery loop, steam economy 2.38.", "AutoCAD 2025 + Aspen Plus\nISO A2 (monochrome.ctb)"),
    ("Annexure H", "DWG-FIN-15T-011\nElectrical Single Line Diagram & SCADA", "11kV / 415V 500kVA TR-01, 250kVA DG, 11 MCC feeders, ISA-95 SCADA Profinet ring topology.", "AutoCAD 2025 (Vector PDF)\nISO A2 (monochrome.ctb)"),
    ("Annexure I", "DWG-FIN-15T-014\nPlant GA Master Assembly Drawing", "SolidWorks 3D B-Rep assembly (PLANT-15T, 12 systems, 27.08 tonnes), isometric & orthographic GAs.", "SolidWorks 2025 SP01.2\nA1 Master Assembly Drawing"),
    ("Master Book", "FINNOAQ_15TPD_CONCEPTUAL_FEED_DRAWINGS_REV_V8.pdf", "Consolidated 8-sheet conceptual FEED package compiled into a single publication-grade book.", "AutoCAD 2025 Native Plot\n625,497 Bytes (Verified)"),
]

for r_idx, row_data in enumerate(anx_rows):
    row = anx_tbl.rows[r_idx + 1]
    for c_idx, val in enumerate(row_data):
        row.cells[c_idx].paragraphs[0].add_run(val)
style_table(anx_tbl, [1.0, 2.27, 2.2, 1.2], [WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.LEFT])

p_disc = doc.add_paragraph(
    "Engineering Provenance Note: All attached drawings strictly observe engineering integrity under FinnoaQ Governance Law. "
    "Drawings originate from native AutoCAD accoreconsole plotting and SolidWorks B-Rep models with 0% synthetic approximations. "
    "All plant engineering deliverables remain conceptual FEED baselines under Gate G-A hold, pending empirical Phase 1 validation."
)
p_disc.runs[0].italic = True
format_para(p_disc, space_before=4, space_after=8)

# ==============================================================================
# 9. ACCEPTANCE & PROJECT AUTHORIZATION FORM
# ==============================================================================
add_header("9. Project Acceptance & Start Authorization Form")

p11 = doc.add_paragraph(
    "To confirm acceptance of this proposal and authorize commencement of Phase 1, please complete and sign below, or issue a formal "
    "Purchase Order / Work Order referencing Document Ref: FNQ-PR-2026-001 (Rev 8.0):"
)
format_para(p11, space_before=2, space_after=6)

sign_tbl = doc.add_table(rows=6, cols=2)
sign_data = [
    ("Client Corporate Entity Name:", ""),
    ("Authorized Representative Name & Title:", ""),
    ("Corporate Registered Address & GSTIN:", ""),
    ("Authorized Signature & Corporate Seal:", "\n\n_________________________________________"),
    ("Execution Date:", ""),
    ("Target Sample Dispatch Date:", "")
]
for i, (k, v) in enumerate(sign_data):
    r = sign_tbl.rows[i]
    r.cells[0].paragraphs[0].add_run(k).bold = True
    r.cells[1].paragraphs[0].add_run(v)
style_table(sign_tbl, [2.5, 4.17], [WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.LEFT])

p_footer = doc.add_paragraph()
r_f = p_footer.add_run("\nFINNOAQ AGRO-INDUSTRIAL TECHNOLOGIES LTD. | CONFIDENTIAL CLIENT ONBOARDING DOSSIER")
r_f.font.size = Pt(8)
r_f.font.color.rgb = CLR_MUTED
p_footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
format_para(p_footer, space_before=8, space_after=0)

# Save DOCX
doc.save(DOCX_OUT)
print(f"Successfully saved: {DOCX_OUT} ({os.path.getsize(DOCX_OUT):,} bytes)", flush=True)

# Export to PDF via Word COM Automation
print("Exporting to PDF via Word COM Automation...", flush=True)
try:
    word = win32com.client.Dispatch("Word.Application")
    word.Visible = False
    doc_word = word.Documents.Open(DOCX_OUT)
    # wdFormatPDF = 17
    doc_word.SaveAs(PDF_OUT, FileFormat=17)
    doc_word.Close()
    word.Quit()
    print(f"Successfully generated client PDF: {PDF_OUT} ({os.path.getsize(PDF_OUT):,} bytes)", flush=True)
except Exception as e:
    print(f"Word COM PDF export notice/error: {e}", flush=True)

print("V8 Client Proposal Generation Complete!", flush=True)
